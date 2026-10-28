import pandas as pd
import openpyxl
import os
import io
import re
import csv
from datetime import datetime, timedelta

THAI_MONTHS = ['มกราคม', 'กุมภาพันธ์', 'มีนาคม', 'เมษายน', 'พฤษภาคม', 'มิถุนายน', 'กรกฎาคม', 'สิงหาคม', 'กันยายน', 'ตุลาคม', 'พฤศจิกายน', 'ธันวาคม']
THAI_MONTHS_MAP = {
    'ม.ค.': 1, 'มกราคม': 1, 'ก.พ.': 2, 'กุมภาพันธ์': 2, 'มี.ค.': 3, 'มีนาคม': 3,
    'เม.ย.': 4, 'เมษายน': 4, 'พ.ค.': 5, 'พฤษภาคม': 5, 'มิ.ย.': 6, 'มิถุนายน': 6,
    'ก.ค.': 7, 'กรกฎาคม': 7, 'ส.ค.': 8, 'สิงหาคม': 8, 'ก.ย.': 9, 'กันยายน': 9,
    'ต.ค.': 10, 'ตุลาคม': 10, 'พ.ย.': 11, 'พฤศจิกายน': 11, 'ธ.ค.': 12, 'ธันวาคม': 12
}

TYPE_CODE_MAP = {
    'ป': 'ลาป่วย',
    'ก': 'ลากิจ',
    'พ': 'ลาพักผ่อน',
    'ส': 'สาย',
    'ข': 'ขาดราชการ',
    'ชภ': 'ลาช่วยภริยาคลอดบุตร',
    'ปภ': 'ลาช่วยภริยาคลอดบุตร',
    'ค': 'ลาคลอดบุตร',
    'อ': 'ลาอื่นๆ',
}

def parse_date(s: str):
    if not s:
        return None
    s = s.strip()
    m1 = re.match(r'^(\d{1,2})[\/\-](\d{1,2})[\/\-](\d{4})$', s)
    if m1:
        d, m, y = int(m1.group(1)), int(m1.group(2)), int(m1.group(3))
        if y > 2400:
            y -= 543
        return datetime(y, m, d)
    m2 = re.match(r'^(\d{1,2})\s*([^\d\s]+)\s*(\d{4})$', s)
    if m2:
        d = int(m2.group(1))
        m_str = m2.group(2).strip()
        y = int(m2.group(3))
        if y > 2400:
            y -= 543
        m = THAI_MONTHS_MAP.get(m_str, 10)
        return datetime(y, m, d)
    return None

def parse_official_report_data(rows_data, header_meta=None):
    """แปลงข้อมูลแถวจากรายงานทางการ (เช่น ข้าราชการ.xlsx / .csv) เป็น (df_details, df_summary)"""
    if header_meta is None:
        header_meta = {}

    details = []
    summary_records = []

    for r_vals in rows_data:
        if not r_vals or len(r_vals) < 13:
            continue
        c1 = r_vals[0]
        if not c1 or pd.isna(c1):
            continue
        c1_str = str(c1).strip()
        if 'ชื่อ' in c1_str or 'กรม' in c1_str or 'รายงาน' in c1_str or 'ตั้งแต่' in c1_str or 'รูปแบบ' in c1_str:
            continue

        name = re.sub(r'^\d+[\.\s]+', '', c1_str).strip()
        sick_days = r_vals[1] if pd.notna(r_vals[1]) and r_vals[1] != '' else None
        sick_times = r_vals[2] if pd.notna(r_vals[2]) and r_vals[2] != '' else None
        bus_days = r_vals[3] if pd.notna(r_vals[3]) and r_vals[3] != '' else None
        bus_times = r_vals[4] if pd.notna(r_vals[4]) and r_vals[4] != '' else None
        tot_sb_days = r_vals[5] if pd.notna(r_vals[5]) and r_vals[5] != '' else None
        tot_sb_times = r_vals[6] if pd.notna(r_vals[6]) and r_vals[6] != '' else None
        late_days = r_vals[7] if pd.notna(r_vals[7]) and r_vals[7] != '' else None
        vac_days = r_vals[8] if pd.notna(r_vals[8]) and r_vals[8] != '' else None
        vac_times = r_vals[9] if pd.notna(r_vals[9]) and r_vals[9] != '' else None
        other_days = r_vals[10] if pd.notna(r_vals[10]) and r_vals[10] != '' else None
        absent_days = r_vals[11] if pd.notna(r_vals[11]) and r_vals[11] != '' else None
        remarks = str(r_vals[12]).strip() if pd.notna(r_vals[12]) and str(r_vals[12]).strip() != 'None' else ''

        summary_records.append({
            'ชื่อ - สกุล': c1_str,
            'Name': name,
            'ลาป่วย_วันทำการ': sick_days,
            'ลาป่วย_จำนวนครั้ง': sick_times,
            'ลากิจ_วันทำการ': bus_days,
            'ลากิจ_จำนวนครั้ง': bus_times,
            'รวมลาป่วย_กิจ_วันทำการ': tot_sb_days,
            'รวมลาป่วย_กิจ_จำนวนครั้ง': tot_sb_times,
            'สาย_วันทำการ': late_days,
            'ลาพักผ่อน_วันทำการ': vac_days,
            'ลาพักผ่อน_จำนวนครั้ง': vac_times,
            'ลาอื่นๆ_วันทำการ': other_days,
            'ขาดราชการ_วันทำการ': absent_days,
            'หมายเหตุ': remarks
        })

        brackets = re.findall(r'([ก-ฮ]+)\[(.*?)\]', remarks)
        for code, span in brackets:
            l_type = TYPE_CODE_MAP.get(code, f'ลา{code}')
            is_half = False
            if '0.5' in remarks or 'ครึ่ง' in remarks or '.5' in remarks:
                if l_type == 'ลาป่วย' and (sick_days % 1 != 0):
                    is_half = True
                elif l_type == 'ลากิจ' and (bus_days % 1 != 0):
                    is_half = True
                elif l_type == 'ลาพักผ่อน' and (vac_days % 1 != 0):
                    is_half = True
                elif ('0.5 วัน' in remarks or '.5 วัน' in remarks) and len(brackets) == 1:
                    is_half = True

            span = span.strip()
            if '-' in span:
                parts = span.split('-')
                d1 = parse_date(parts[0])
                d2 = parse_date(parts[1])
                if d1 and d2:
                    cur = d1
                    while cur <= d2:
                        m_th = THAI_MONTHS[cur.month - 1]
                        b_year = str(cur.year + 543 if cur.month < 10 else cur.year + 544)
                        details.append({
                            'Name': name,
                            'Month': m_th,
                            'Day': cur.day,
                            'Type': l_type,
                            'BudgetYear': b_year,
                            'Remarks': remarks,
                            'IsHalf': is_half,
                            'Date': cur
                        })
                        cur += timedelta(days=1)
            else:
                d = parse_date(span)
                if d:
                    m_th = THAI_MONTHS[d.month - 1]
                    b_year = str(d.year + 543 if d.month < 10 else d.year + 544)
                    details.append({
                        'Name': name,
                        'Month': m_th,
                        'Day': d.day,
                        'Type': l_type,
                        'BudgetYear': b_year,
                        'Remarks': remarks,
                        'IsHalf': is_half,
                        'Date': d
                    })

    df_details = pd.DataFrame(details)
    df_summary = pd.DataFrame(summary_records)
    return df_details, df_summary

def load_data(file_source):
    """
    ฟังก์ชันโหลดข้อมูลตารางวันลา รองรับทั้ง 2 รูปแบบอัตโนมัติ:
    1. รูปแบบรายงานทางการกรมเจ้าท่า (ข้าราชการ.xlsx / .csv)
    2. รูปแบบตารางตารางวันลาหลายปีงบประมาณ (ต.ค 67-ก.ย 68, ต.ค 68-ก.ย 69)
    """
    try:
        # Handle file-like objects (like Streamlit UploadedFile) or string paths
        fname = getattr(file_source, 'name', str(file_source)).lower()
        if hasattr(file_source, 'read'):
            file_content = file_source.read()
            file_source.seek(0)
        else:
            if not os.path.exists(file_source):
                return pd.DataFrame()
            with open(file_source, 'rb') as f:
                file_content = f.read()

        # ตรวจสอบว่าเป็นไฟล์ CSV หรือไม่
        if fname.endswith('.csv'):
            try:
                text_content = file_content.decode('utf-8-sig', errors='replace')
            except Exception:
                text_content = file_content.decode('cp874', errors='replace')
            reader = csv.reader(io.StringIO(text_content))
            rows = [row for row in reader]

            header_meta = {}
            if len(rows) > 0 and rows[0]: header_meta['org'] = rows[0][0]
            if len(rows) > 1 and rows[1]: header_meta['title'] = rows[1][0]
            if len(rows) > 2 and rows[2]: header_meta['date_range'] = rows[2][0]
            if len(rows) > 3 and rows[3]: header_meta['format'] = rows[3][0]

            df_details, df_summary = parse_official_report_data(rows, header_meta)
            df_details.attrs['official_summary'] = df_summary
            df_details.attrs['header_meta'] = header_meta
            return df_details

        # ลองเปิดด้วย Openpyxl
        source_for_ox = io.BytesIO(file_content)
        wb = openpyxl.load_workbook(source_for_ox, data_only=True)
        sheets = wb.sheetnames

        # ตรวจสอบว่าเป็นแบบฟอร์มทางการ (รายงานการมาปฏิบัติราชการ) หรือไม่
        is_official = False
        target_official_sheet = None
        for s in sheets:
            if s == 'รายงานการมาปฏิบัติราชการ':
                is_official = True
                target_official_sheet = s
                break
            ws_chk = wb[s]
            sample = ' '.join(str(ws_chk.cell(row=r, column=1).value or '') for r in range(1, min(6, ws_chk.max_row + 1)))
            if 'รายงานการมาปฏิบัติราชการ' in sample:
                is_official = True
                target_official_sheet = s
                break

        if is_official and target_official_sheet:
            ws = wb[target_official_sheet]
            rows = []
            for r in range(1, ws.max_row + 1):
                rows.append([ws.cell(row=r, column=c).value for c in range(1, 14)])

            header_meta = {
                'org': ws.cell(row=1, column=1).value or 'กรมเจ้าท่า',
                'title': ws.cell(row=2, column=1).value or 'รายงานการมาปฏิบัติราชการของข้าราชการพลเรือนสามัญ',
                'date_range': ws.cell(row=3, column=1).value or 'ตั้งแต่วันที่ 01/10/2569 ถึงวันที่ 28/10/2569',
                'format': ws.cell(row=4, column=1).value or 'รูปแบบการออกรายงาน : โครงสร้างตามกฎหมาย - ทั้งส่วนราชการ - กรมเจ้าท่า - กระทรวงคมนาคม',
            }
            df_details, df_summary = parse_official_report_data(rows, header_meta)
            df_details.attrs['official_summary'] = df_summary
            df_details.attrs['header_meta'] = header_meta
            return df_details

        # ตรวจสอบแบบตารางวันลาเดิม (Matrix: ต.ค 67-ก.ย 68)
        source_for_pd = io.BytesIO(file_content)
        excel_file = pd.ExcelFile(source_for_pd)
        relevant_sheets = [s for s in sheets if re.search(r'ต\.?ค.*ก\.?ย', s)]
        if not relevant_sheets:
            relevant_sheets = [s for s in ['ต.ค 67-ก.ย 68', 'ต.ค 68-ก.ย 69 (2)'] if s in sheets]

        all_details = []
        for sheet_name in relevant_sheets:
            if sheet_name not in sheets:
                continue

            df = pd.read_excel(source_for_pd, sheet_name=sheet_name)
            source_for_pd.seek(0)

            month_col = 'เดือน'
            if month_col not in df.columns:
                continue

            ws = wb[sheet_name]
            name_col = df.columns[0]
            day_cols = df.columns[2:33]
            remark_col = 'หมายเหตุ'

            year_match = re.search(r'ก\.?ย\.?\s*(\d{2,4})', sheet_name)
            if year_match:
                y_val = int(year_match.group(1))
                budget_year = str(2500 + y_val if y_val < 100 else y_val)
            else:
                budget_year = '2568' if '67' in sheet_name else '2569'

            type_map = {'ป': 'ลาป่วย', 'ก': 'ลากิจ', 'พ': 'ลาพักผ่อน', 'ส': 'สาย', 'ป.': 'ลาป่วย', 'ข': 'ขาดราชการ'}

            for idx, row in df.iterrows():
                name = row[name_col]
                month = row[month_col]

                if pd.isna(month) or month in ['Aเดือน', 'รวม', 'เดือน']:
                    continue
                if pd.isna(name) or 'Aวันที่' in str(name):
                    continue

                remarks = str(row[remark_col]) if pd.notna(row[remark_col]) else ""

                for day_idx, day_val in enumerate(day_cols):
                    val = row[day_val]
                    if pd.notna(val) and str(val).strip() in type_map:
                        leave_type = type_map[str(val).strip()]

                        excel_row = idx + 2
                        excel_col = day_idx + 3
                        cell = ws.cell(row=excel_row, column=excel_col)
                        is_half = False

                        fill = cell.fill
                        if fill and fill.start_color and getattr(fill.start_color, 'rgb', None):
                            color = str(fill.start_color.rgb)
                            if color == 'FFFFFF00' or 'FFFF00' in color:
                                is_half = True

                        day_remark = remarks
                        if 'ครึ่ง' in remarks and not is_half:
                            day_remark = ""

                        all_details.append({
                            'Name': name,
                            'Month': month,
                            'Day': day_idx + 1,
                            'Type': leave_type,
                            'BudgetYear': budget_year,
                            'Remarks': day_remark,
                            'IsHalf': is_half
                        })

        return pd.DataFrame(all_details)
    except Exception as e:
        print(f"Error loading file: {e}")
        return pd.DataFrame()

def generate_official_leave_excel(df_summary: pd.DataFrame, header_meta: dict = None) -> bytes:
    """
    ส่งออกไฟล์ Excel รายงานราชการตามแบบฟอร์มทางการของกรมเจ้าท่า
    (ฟอร์แมต 13 คอลัมน์ ฟอนต์ TH SarabunPSK 16pt และการผสานเซลล์ A1:M1, A5:A6, B5:C5 ฯลฯ ตรงตามต้นฉบับ 100%)
    """
    if header_meta is None:
        header_meta = {}

    org_name = header_meta.get('org') or 'กรมเจ้าท่า'
    title = header_meta.get('title') or 'รายงานการมาปฏิบัติราชการของข้าราชการพลเรือนสามัญ'
    date_range = header_meta.get('date_range') or 'ตั้งแต่วันที่ 01/10/2569 ถึงวันที่ 28/10/2569'
    report_format = header_meta.get('format') or 'รูปแบบการออกรายงาน : โครงสร้างตามกฎหมาย - ทั้งส่วนราชการ - กรมเจ้าท่า - กระทรวงคมนาคม'

    buffer = io.BytesIO()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "รายงานการมาปฏิบัติราชการ"

    # Fonts
    font_title = openpyxl.styles.Font(name='TH SarabunPSK', size=16, bold=True)
    font_hdr = openpyxl.styles.Font(name='TH SarabunPSK', size=16, bold=True)
    font_data = openpyxl.styles.Font(name='TH SarabunPSK', size=16, bold=False)

    # Alignments
    align_center = openpyxl.styles.Alignment(horizontal='center', vertical='center', wrap_text=True)
    align_left = openpyxl.styles.Alignment(horizontal='left', vertical='center', wrap_text=True)

    # Borders
    border_thin = openpyxl.styles.Border(
        left=openpyxl.styles.Side(border_style='thin', color='000000'),
        right=openpyxl.styles.Side(border_style='thin', color='000000'),
        top=openpyxl.styles.Side(border_style='thin', color='000000'),
        bottom=openpyxl.styles.Side(border_style='thin', color='000000')
    )

    # Row heights for header
    for r in range(1, 7):
        ws.row_dimensions[r].height = 21.0

    # Titles (Rows 1-4)
    ws.merge_cells('A1:M1')
    ws['A1'] = org_name
    ws['A1'].font = font_title
    ws['A1'].alignment = openpyxl.styles.Alignment(horizontal='center')

    ws.merge_cells('A2:M2')
    ws['A2'] = title
    ws['A2'].font = font_title
    ws['A2'].alignment = openpyxl.styles.Alignment(horizontal='center')

    ws.merge_cells('A3:M3')
    ws['A3'] = date_range
    ws['A3'].font = font_title
    ws['A3'].alignment = openpyxl.styles.Alignment(horizontal='center')

    ws.merge_cells('A4:M4')
    ws['A4'] = report_format
    ws['A4'].font = font_title
    ws['A4'].alignment = openpyxl.styles.Alignment(horizontal='left')

    # Headers (Rows 5-6)
    ws.merge_cells('A5:A6'); ws['A5'] = 'ชื่อ - สกุล'
    ws.merge_cells('B5:C5'); ws['B5'] = 'ลาป่วย'
    ws.merge_cells('D5:E5'); ws['D5'] = 'ลากิจ'
    ws.merge_cells('F5:G5'); ws['F5'] = 'รวมลาป่วย กิจ'
    ws.merge_cells('H5:H6'); ws['H5'] = 'สาย\n(วันทำการ)'
    ws.merge_cells('I5:J5'); ws['I5'] = 'ลาพักผ่อน'
    ws.merge_cells('K5:K6'); ws['K5'] = 'ลาอื่นๆ\n(วันทำการ)'
    ws.merge_cells('L5:L6'); ws['L5'] = 'ขาดราชการ\n(วันทำการ)'
    ws.merge_cells('M5:M6'); ws['M5'] = 'หมายเหตุ'

    ws['B6'] = 'วันทำการ'; ws['C6'] = 'จำนวนครั้ง'
    ws['D6'] = 'วันทำการ'; ws['E6'] = 'จำนวนครั้ง'
    ws['F6'] = 'วันทำการ'; ws['G6'] = 'จำนวนครั้ง'
    ws['I6'] = 'วันทำการ'; ws['J6'] = 'จำนวนครั้ง'

    for r in [5, 6]:
        for c in range(1, 14):
            cell = ws.cell(row=r, column=c)
            cell.font = font_hdr
            cell.alignment = align_center
            cell.border = border_thin

    # Column widths matching original template
    ws.column_dimensions['A'].width = 30.0
    ws.column_dimensions['B'].width = 15.0
    for col_c in ['C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'K', 'L']:
        ws.column_dimensions[col_c].width = 13.0
    ws.column_dimensions['M'].width = 35.0

    # Data rows (Row 7+)
    row_idx = 7
    for _, r_data in df_summary.iterrows():
        ws.row_dimensions[row_idx].height = 30.0

        c1_val = r_data.get('ชื่อ - สกุล') or r_data.get('Name') or ''
        c1 = ws.cell(row=row_idx, column=1, value=c1_val)
        c1.font = font_data
        c1.alignment = align_left
        c1.border = border_thin

        num_keys = [
            'ลาป่วย_วันทำการ', 'ลาป่วย_จำนวนครั้ง',
            'ลากิจ_วันทำการ', 'ลากิจ_จำนวนครั้ง',
            'รวมลาป่วย_กิจ_วันทำการ', 'รวมลาป่วย_กิจ_จำนวนครั้ง',
            'สาย_วันทำการ',
            'ลาพักผ่อน_วันทำการ', 'ลาพักผ่อน_จำนวนครั้ง',
            'ลาอื่นๆ_วันทำการ', 'ขาดราชการ_วันทำการ'
        ]
        for c_i, k in enumerate(num_keys, start=2):
            val = r_data.get(k)
            if pd.isna(val) or val is None or val == '':
                cell_val = None
            else:
                try:
                    f_val = float(val)
                    cell_val = int(f_val) if f_val.is_integer() else f_val
                except Exception:
                    cell_val = val
            c_cell = ws.cell(row=row_idx, column=c_i, value=cell_val)
            c_cell.font = font_data
            c_cell.alignment = align_center
            c_cell.border = border_thin

        rem_val = r_data.get('หมายเหตุ', '')
        if pd.isna(rem_val) or rem_val is None or str(rem_val).strip() == 'None':
            rem_val = None
        c13 = ws.cell(row=row_idx, column=13, value=rem_val)
        c13.font = font_data
        c13.alignment = align_left
        c13.border = border_thin

        row_idx += 1

    wb.save(buffer)
    return buffer.getvalue()

def synthesize_official_summary(df_details, roster=None, year_be=2569, month_num=10, month_name='ตุลาคม'):
    """สร้างตารางสรุปรายงานราชการ 13 คอลัมน์ตามแบบฟอร์มกรมเจ้าท่าจากรายการวันลารายวัน"""
    if roster is None:
        roster = []

    df_m = df_details.copy() if not df_details.empty else pd.DataFrame()
    if not df_m.empty:
        if month_name and 'Month' in df_m.columns:
            df_m = df_m[df_m['Month'] == month_name]
        if year_be and 'BudgetYear' in df_m.columns:
            df_m = df_m[df_m['BudgetYear'].astype(str) == str(year_be)]
        df_m['LeaveVal'] = df_m['IsHalf'].apply(lambda x: 0.5 if x else 1.0) if 'IsHalf' in df_m.columns else 1.0

    if roster:
        names_list = [p['name'] for p in roster]
    elif not df_m.empty:
        names_list = sorted(df_m['Name'].unique())
    else:
        names_list = []

    summary_rows = []
    type_code = {'ลาป่วย': 'ป', 'ลาพักผ่อน': 'พ', 'ลากิจ': 'ก', 'สาย': 'ส', 'ขาดราชการ': 'ข'}

    for idx, name in enumerate(names_list, start=1):
        clean_k = re.sub(r'[\s\.\d]+', '', name)
        p_rows = df_m[df_m['Name'].str.replace(r'[\s\.\d]+', '', regex=True) == clean_k] if not df_m.empty else pd.DataFrame()

        if p_rows.empty:
            summary_rows.append({
                'ชื่อ - สกุล': f'{idx}. {name}',
                'Name': name,
                'ลาป่วย_วันทำการ': None, 'ลาป่วย_จำนวนครั้ง': None,
                'ลากิจ_วันทำการ': None, 'ลากิจ_จำนวนครั้ง': None,
                'รวมลาป่วย_กิจ_วันทำการ': None, 'รวมลาป่วย_กิจ_จำนวนครั้ง': None,
                'สาย_วันทำการ': None,
                'ลาพักผ่อน_วันทำการ': None, 'ลาพักผ่อน_จำนวนครั้ง': None,
                'ลาอื่นๆ_วันทำการ': None, 'ขาดราชการ_วันทำการ': None,
                'หมายเหตุ': None
            })
            continue

        sick_p = p_rows[p_rows['Type'] == 'ลาป่วย']
        sick_days = sick_p['LeaveVal'].sum() if not sick_p.empty else None
        sick_times = len(sick_p.drop_duplicates(subset=['Day'])) if not sick_p.empty else None

        bus_p = p_rows[p_rows['Type'] == 'ลากิจ']
        bus_days = bus_p['LeaveVal'].sum() if not bus_p.empty else None
        bus_times = len(bus_p.drop_duplicates(subset=['Day'])) if not bus_p.empty else None

        tot_sb_days = (sick_days or 0) + (bus_days or 0) if (sick_days or bus_days) else None
        tot_sb_times = (sick_times or 0) + (bus_times or 0) if (sick_times or bus_times) else None

        late_p = p_rows[p_rows['Type'] == 'สาย']
        late_days = len(late_p) if not late_p.empty else None

        vac_p = p_rows[p_rows['Type'] == 'ลาพักผ่อน']
        vac_days = vac_p['LeaveVal'].sum() if not vac_p.empty else None
        vac_times = len(vac_p.drop_duplicates(subset=['Day'])) if not vac_p.empty else None

        other_p = p_rows[~p_rows['Type'].isin(['ลาป่วย', 'ลากิจ', 'ลาพักผ่อน', 'สาย', 'ขาดราชการ'])]
        other_days = other_p['LeaveVal'].sum() if not other_p.empty else None

        absent_p = p_rows[p_rows['Type'] == 'ขาดราชการ']
        absent_days = len(absent_p) if not absent_p.empty else None

        total_working_days = (sick_days or 0) + (bus_days or 0) + (vac_days or 0) + (other_days or 0)
        tot_str = f'{total_working_days:g}'

        spans = []
        for l_t in ['ลาพักผ่อน', 'ลาป่วย', 'ลากิจ']:
            sub = p_rows[p_rows['Type'] == l_t].sort_values('Day')
            if not sub.empty:
                c = type_code.get(l_t, l_t)
                days = sorted(sub['Day'].unique())
                cur_group = []
                for d in days:
                    if not cur_group:
                        cur_group.append(d)
                    elif d == cur_group[-1] + 1:
                        cur_group.append(d)
                    else:
                        if len(cur_group) == 1:
                            spans.append(f'{c}[{cur_group[0]:02d}/{month_num:02d}/{year_be}]')
                        else:
                            spans.append(f'{c}[{cur_group[0]:02d}/{month_num:02d}/{year_be}-{cur_group[-1]:02d}/{month_num:02d}/{year_be}]')
                        cur_group = [d]
                if cur_group:
                    if len(cur_group) == 1:
                        spans.append(f'{c}[{cur_group[0]:02d}/{month_num:02d}/{year_be}]')
                    else:
                        spans.append(f'{c}[{cur_group[0]:02d}/{month_num:02d}/{year_be}-{cur_group[-1]:02d}/{month_num:02d}/{year_be}]')

        remarks = f'วันทำการ {tot_str} วัน ,' + ' ,'.join(spans) if spans else None

        summary_rows.append({
            'ชื่อ - สกุล': f'{idx}. {name}',
            'Name': name,
            'ลาป่วย_วันทำการ': sick_days, 'ลาป่วย_จำนวนครั้ง': sick_times,
            'ลากิจ_วันทำการ': bus_days, 'ลากิจ_จำนวนครั้ง': bus_times,
            'รวมลาป่วย_กิจ_วันทำการ': tot_sb_days, 'รวมลาป่วย_กิจ_จำนวนครั้ง': tot_sb_times,
            'สาย_วันทำการ': late_days,
            'ลาพักผ่อน_วันทำการ': vac_days, 'ลาพักผ่อน_จำนวนครั้ง': vac_times,
            'ลาอื่นๆ_วันทำการ': other_days, 'ขาดราชการ_วันทำการ': absent_days,
            'หมายเหตุ': remarks
        })

    return pd.DataFrame(summary_rows)
